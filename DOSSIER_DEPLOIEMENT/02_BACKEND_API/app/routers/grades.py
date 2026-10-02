from fastapi import APIRouter, Depends, HTTPException, status, UploadFile
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from typing import List, Optional
from datetime import datetime, date as date_type
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import tempfile
import os
import io

from .. import crud, schemas, models
from ..database import get_db
from ..routers.auth import get_school_scope, SchoolScope

router = APIRouter(tags=["Grades, Attendance & Bulletins"])

# ----------------- EVALUATIONS / GRADES -----------------
@router.get("/evaluations", response_model=List[schemas.EvaluationResponse])
def read_evaluations(
    classe_id: Optional[int] = None, 
    eleve_id: Optional[int] = None, 
    devoir_numero: Optional[str] = None,
    db: Session = Depends(get_db),
    scope: SchoolScope = Depends(get_school_scope)
):
    return crud.get_evaluations(
        db,
        classe_id=classe_id,
        eleve_id=eleve_id,
        devoir_numero=devoir_numero,
        ecole_id=scope.ecole_id,
        code_etablissement=scope.code_etablissement
    )

class EvaluationGradeItem(schemas.BaseModel):
    eleve_id: int
    note: float
    appreciation: Optional[str] = None
    eval_id: Optional[int] = None
    justification_modification: Optional[str] = None
    autorisation_modification: Optional[bool] = False

class BulkEvaluationSave(schemas.BaseModel):
    classe_id: int
    matiere: str
    trimestre: int = 1
    type: str = "interrogation écrite"
    devoir_numero: Optional[str] = None
    coefficient: int = 1
    date: date_type
    date_recuperation: Optional[date_type] = None
    date_correction: Optional[date_type] = None
    date_remise: Optional[date_type] = None
    valide: bool = False
    notes: List[EvaluationGradeItem]

@router.post("/evaluations/bulk")
def save_evaluations_bulk(
    payload: BulkEvaluationSave,
    db: Session = Depends(get_db),
    scope: SchoolScope = Depends(get_school_scope)
):
    # ─── Vérification du verrou de période ────────────────────────────
    today = date_type.today()
    periode = db.query(models.PeriodeScolaire).filter(
        models.PeriodeScolaire.numero == payload.trimestre,
        models.PeriodeScolaire.type_periode == "trimestre"
    ).first()

    if periode:
        cloture_atteinte = (today > periode.date_cloture_saisie) if periode.date_cloture_saisie else False
        est_bloque = (cloture_atteinte or bool(periode.cloture_forcee)) and not bool(periode.deverrouille_par_admin)
        if est_bloque:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail=f"La période de saisie des notes pour le « {periode.libelle} » est clôturée depuis le {periode.date_cloture_saisie}. Seul un administrateur peut sauter ce verrou."
            )

    saved_evals = []
    code_etab = scope.code_etablissement
    if not code_etab and payload.classe_id:
        cls = db.query(models.Classe).filter(models.Classe.id == payload.classe_id).first()
        if cls and cls.ET_CODEETABLISSEMENT:
            code_etab = cls.ET_CODEETABLISSEMENT

    for item in payload.notes:
        db_eval = None
        if item.eval_id:
            db_eval = db.query(models.Evaluation).filter(models.Evaluation.id == item.eval_id).first()
        elif payload.devoir_numero:
            db_eval = db.query(models.Evaluation).filter(
                models.Evaluation.classe_id == payload.classe_id,
                models.Evaluation.eleve_id == item.eleve_id,
                models.Evaluation.devoir_numero == payload.devoir_numero
            ).first()

        if db_eval:
            # Update existing
            if db_eval.valide and not item.autorisation_modification:
                continue
            db_eval.matiere = payload.matiere
            db_eval.type = payload.type
            db_eval.trimestre = payload.trimestre
            db_eval.note = float(item.note)
            db_eval.coefficient = payload.coefficient
            db_eval.date = payload.date
            db_eval.appreciation = item.appreciation
            db_eval.date_recuperation = payload.date_recuperation
            db_eval.date_correction = payload.date_correction
            db_eval.date_remise = payload.date_remise
            db_eval.valide = payload.valide
            if item.justification_modification:
                db_eval.justification_modification = item.justification_modification
        else:
            # Create new
            db_eval = models.Evaluation(
                matiere=payload.matiere,
                type=payload.type,
                trimestre=payload.trimestre,
                note=float(item.note),
                coefficient=payload.coefficient,
                date=payload.date,
                appreciation=item.appreciation,
                classe_id=payload.classe_id,
                eleve_id=item.eleve_id,
                devoir_numero=payload.devoir_numero,
                date_recuperation=payload.date_recuperation,
                date_correction=payload.date_correction,
                date_remise=payload.date_remise,
                valide=payload.valide,
                ET_CODEETABLISSEMENT=code_etab,
                date_creation=datetime.utcnow()
            )
            db.add(db_eval)

        saved_evals.append(db_eval)

    db.commit()
    return {"success": True, "count": len(saved_evals), "message": f"{len(saved_evals)} notes enregistrées avec succès"}



# ----------------- ATTENDANCE / PRESENCE -----------------
@router.get("/attendance", response_model=List[schemas.PresenceResponse])
def read_presences(
    classe_id: Optional[int] = None, 
    eleve_id: Optional[int] = None, 
    date: Optional[str] = None,
    db: Session = Depends(get_db),
    scope: SchoolScope = Depends(get_school_scope)
):
    from datetime import datetime
    date_val = None
    if date:
        try:
            date_val = datetime.strptime(date, "%Y-%m-%d").date()
        except Exception:
            pass
    return crud.get_presences(
        db,
        classe_id=classe_id,
        eleve_id=eleve_id,
        date_val=date_val,
        ecole_id=scope.ecole_id,
        code_etablissement=scope.code_etablissement
    )

@router.get("/attendance/{presence_id}", response_model=schemas.PresenceResponse)
def read_presence(
    presence_id: int,
    db: Session = Depends(get_db),
    scope: SchoolScope = Depends(get_school_scope)
):
    db_presence = db.query(models.Presence).filter(models.Presence.id == presence_id).first()
    if not db_presence:
        raise HTTPException(status_code=404, detail="Présence introuvable")
    
    if not scope.is_global and scope.code_etablissement:
        if db_presence.ET_CODEETABLISSEMENT and db_presence.ET_CODEETABLISSEMENT != scope.code_etablissement:
            raise HTTPException(status_code=403, detail="Accès refusé à cette présence d'un autre établissement")

    # Add locked status
    time_elapsed = (datetime.utcnow() - db_presence.date_creation).total_seconds()
    db_presence.locked = time_elapsed > 3600
    return db_presence

@router.post("/attendance", response_model=schemas.PresenceResponse, status_code=status.HTTP_201_CREATED)
def create_new_presence(
    presence: schemas.PresenceCreate,
    db: Session = Depends(get_db),
    scope: SchoolScope = Depends(get_school_scope)
):
    code_etab = scope.code_etablissement if not scope.is_global else None
    return crud.create_presence(db=db, presence=presence, code_etablissement=code_etab, ecole_id=scope.ecole_id)

@router.put("/attendance/{presence_id}", response_model=schemas.PresenceResponse)
def update_existing_presence(
    presence_id: int,
    presence_in: schemas.PresenceBase,
    db: Session = Depends(get_db),
    scope: SchoolScope = Depends(get_school_scope)
):
    db_presence = db.query(models.Presence).filter(models.Presence.id == presence_id).first()
    if not db_presence:
        raise HTTPException(status_code=404, detail="Présence introuvable")
    
    if not scope.is_global and scope.code_etablissement:
        if db_presence.ET_CODEETABLISSEMENT and db_presence.ET_CODEETABLISSEMENT != scope.code_etablissement:
            raise HTTPException(status_code=403, detail="Accès refusé à cette présence d'un autre établissement")

    time_elapsed = (datetime.utcnow() - db_presence.date_creation).total_seconds()
    if time_elapsed > 3600:
        raise HTTPException(
            status_code=403,
            detail="Appel verrouillé. Le délai d'édition autorisé (1 heure) a expiré."
        )
        
    db_presence.statut = presence_in.statut
    if presence_in.justification:
        db_presence.justification = presence_in.justification
    if presence_in.heure:
        db_presence.heure = presence_in.heure
        
    db.commit()
    db.refresh(db_presence)
    db_presence.locked = time_elapsed > 3600
    return db_presence


# ----------------- BULLETINS -----------------
@router.get("/bulletins")
def read_bulletins(
    classe_id: Optional[int] = None,
    db: Session = Depends(get_db),
    scope: SchoolScope = Depends(get_school_scope)
):
    query = db.query(models.Bulletin)
    if classe_id is not None:
        query = query.filter(models.Bulletin.classe_id == classe_id)
    if scope.code_etablissement:
        query = query.filter(models.Bulletin.classe_id.in_(
            db.query(models.Classe.id).filter(models.Classe.ET_CODEETABLISSEMENT == scope.code_etablissement)
        ))
    elif scope.ecole_id is not None:
        query = query.filter(models.Bulletin.classe_id.in_(
            db.query(models.Classe.id).filter(models.Classe.ecole_id == scope.ecole_id)
        ))
    return query.all()

@router.post("/bulletins/generate")
def generate_bulletin(
    classe_id: int,
    trimestre: int,
    db: Session = Depends(get_db),
    scope: SchoolScope = Depends(get_school_scope)
):
    cls = db.query(models.Classe).filter(models.Classe.id == classe_id).first()
    if not cls:
        raise HTTPException(status_code=404, detail="Classe non trouvée")
    if not scope.is_global and scope.code_etablissement:
        if cls.ET_CODEETABLISSEMENT and cls.ET_CODEETABLISSEMENT != scope.code_etablissement:
            raise HTTPException(status_code=403, detail="Accès refusé à cette classe d'un autre établissement")

    db_bulletin = db.query(models.Bulletin).filter(
        models.Bulletin.classe_id == classe_id,
        models.Bulletin.trimestre == trimestre
    ).first()
    
    if db_bulletin:
        db_bulletin.date_generation = datetime.utcnow()
        db_bulletin.statut = "genere"
        db.commit()
        db.refresh(db_bulletin)
        return {"message": "Bulletin régénéré avec succès", "id": db_bulletin.id}
        
    new_bulletin = models.Bulletin(
        trimestre=trimestre,
        statut="genere",
        date_generation=datetime.utcnow(),
        classe_id=classe_id
    )
    db.add(new_bulletin)
    db.commit()
    db.refresh(new_bulletin)
    return {"message": "Bulletin généré avec succès", "id": new_bulletin.id}


# ----------------- SEANCES / CAHIER DE TEXTE -----------------
@router.get("/seances", response_model=List[schemas.SeanceResponse])
def read_seances(
    classe_id: Optional[int] = None,
    enseignant_id: Optional[int] = None,
    matiere_id: Optional[int] = None,
    db: Session = Depends(get_db),
    scope: SchoolScope = Depends(get_school_scope)
):
    return crud.get_seances(
        db,
        classe_id=classe_id,
        enseignant_id=enseignant_id,
        matiere_id=matiere_id,
        ecole_id=scope.ecole_id,
        code_etablissement=scope.code_etablissement
    )

@router.post("/seances", response_model=schemas.SeanceResponse, status_code=status.HTTP_201_CREATED)
def create_new_seance(
    seance: schemas.SeanceCreate,
    db: Session = Depends(get_db),
    scope: SchoolScope = Depends(get_school_scope)
):
    return crud.create_seance(db, seance, code_etablissement=scope.code_etablissement)


# ----------------- POINTAGES -----------------
@router.get("/pointages", response_model=List[schemas.PointageResponse])
def read_pointages(
    personnel_id: Optional[int] = None,
    date: Optional[str] = None,
    db: Session = Depends(get_db),
    scope: SchoolScope = Depends(get_school_scope)
):
    return crud.get_pointages(db, personnel_id=personnel_id, date_str=date, code_etablissement=scope.code_etablissement)

@router.post("/pointages", response_model=schemas.PointageResponse)
def clock_in_out(
    pointage: schemas.PointageCreate,
    db: Session = Depends(get_db),
    scope: SchoolScope = Depends(get_school_scope)
):
    return crud.create_or_update_pointage(db, pointage, code_etablissement=scope.code_etablissement)


# ----------------- EXCEL EXPORT/IMPORT FOR GRADES -----------------
@router.get("/evaluations/export-template/{classe_id}")
def export_grades_template(classe_id: int, db: Session = Depends(get_db)):
    """
    Exporte un template Excel pour la saisie des notes d'une classe.
    Le fichier contient la liste des élèves avec des colonnes vides pour les notes.
    """
    # Récupérer la classe et ses élèves
    db_class = db.query(models.Classe).filter(models.Classe.id == classe_id).first()
    if not db_class:
        raise HTTPException(status_code=404, detail="Classe non trouvée")
    
    students = db.query(models.Eleve).filter(models.Eleve.classe_id == classe_id).all()
    if not students:
        raise HTTPException(status_code=404, detail="Aucun élève trouvé dans cette classe")
    
    # Créer le workbook Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Saisie des Notes"
    
    # Styles
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # En-têtes
    headers = [
        "ID Élève",
        "Matricule",
        "Nom",
        "Prénom",
        "Classe ID",
        "Matière",
        "Type Devoir",
        "Trimestre",
        "Note (/20)",
        "Coefficient",
        "Date Devoir",
        "Appréciation",
        "Numéro Devoir"
    ]
    
    # Écrire les en-têtes
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border
    
    # Ajuster la largeur des colonnes
    column_widths = [10, 15, 20, 20, 10, 15, 15, 10, 12, 12, 15, 25, 15]
    for col_num, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(col_num)].width = width
    
    # Remplir avec les données des élèves
    for row_num, student in enumerate(students, 2):
        ws.cell(row=row_num, column=1, value=student.id)
        ws.cell(row=row_num, column=2, value=student.matricule)
        ws.cell(row=row_num, column=3, value=student.nom)
        ws.cell(row=row_num, column=4, value=student.prenom)
        ws.cell(row=row_num, column=5, value=student.classe_id)
        
        # Appliquer les bordures aux cellules de données
        for col_num in range(1, 14):
            cell = ws.cell(row=row_num, column=col_num)
            cell.border = border
    
    # Ajouter une ligne d'instructions
    ws.append([])
    ws.append(["INSTRUCTIONS:"])
    ws.append(["- Remplissez TOUTES les colonnes obligatoires pour chaque ligne:"])
    ws.append(["  Matière, Type Devoir, Trimestre, Note, Coefficient, Date Devoir"])
    ws.append(["- Les notes doivent être sur 20"])
    ws.append(["- Types de devoir possibles: devoir, composition, examen, interrogation"])
    ws.append(["- Trimestre: 1, 2 ou 3"])
    ws.append(["- Date format: YYYY-MM-DD (ex: 2026-07-09) ou DD/MM/YYYY (ex: 09/07/2026)"])
    ws.append(["- Coefficient: nombre entier positif (ex: 1, 2, 3, 4)"])
    ws.append(["- Appréciation et Numéro Devoir sont optionnels"])
    
    # Sauvegarder dans un fichier temporaire
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
    wb.save(temp_file.name)
    temp_file.close()
    
    # Nom du fichier
    filename = f"template_notes_{db_class.CE_LIBELLE.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    
    return FileResponse(
        temp_file.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=filename
    )


@router.post("/evaluations/import-excel")
async def import_grades_from_excel(
    file: UploadFile,
    db: Session = Depends(get_db)
):
    """
    Importe les notes depuis un fichier Excel.
    Le fichier doit contenir les colonnes: ID Élève, Matière, Type Devoir, Trimestre, Note, Coefficient, Date Devoir, (optionnel: Appréciation, Numéro Devoir)
    Toutes les données sont lues depuis le fichier Excel.
    """
    if not file or not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="Le fichier doit être un fichier Excel (.xlsx ou .xls)")
    
    try:
        # Lire le fichier Excel
        import openpyxl
        contents = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(contents))
        ws = wb.active
        
        # Lire les en-têtes pour identifier les colonnes
        headers = [cell.value for cell in ws[1]]
        
        # Mapper les noms de colonnes aux indices
        col_map = {}
        for idx, header in enumerate(headers):
            if header:
                header_lower = str(header).lower()
                if 'id' in header_lower and 'élève' in header_lower:
                    col_map['student_id'] = idx
                elif 'matière' in header_lower or 'matiere' in header_lower:
                    col_map['matiere'] = idx
                elif 'type' in header_lower and 'devoir' in header_lower:
                    col_map['type_devoir'] = idx
                elif 'trimestre' in header_lower:
                    col_map['trimestre'] = idx
                elif 'note' in header_lower:
                    col_map['note'] = idx
                elif 'coef' in header_lower or 'coefficient' in header_lower:
                    col_map['coefficient'] = idx
                elif 'date' in header_lower and 'devoir' in header_lower:
                    col_map['date_devoir'] = idx
                elif 'appréciation' in header_lower or 'appreciation' in header_lower:
                    col_map['appreciation'] = idx
                elif 'numéro' in header_lower or 'numero' in header_lower or 'devoir' in header_lower and 'numéro' not in header_lower:
                    col_map['devoir_numero'] = idx
                elif 'classe' in header_lower and 'id' in header_lower:
                    col_map['classe_id'] = idx
        
        # Vérifier les colonnes obligatoires
        required_cols = ['student_id', 'matiere', 'type_devoir', 'trimestre', 'note', 'coefficient', 'date_devoir']
        missing_cols = [col for col in required_cols if col not in col_map]
        if missing_cols:
            raise HTTPException(
                status_code=400, 
                detail=f"Colonnes manquantes dans le fichier Excel: {', '.join(missing_cols)}"
            )
        
        results = {
            "success": 0,
            "errors": 0,
            "error_details": []
        }
        
        # Parcourir les lignes (à partir de la ligne 2)
        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            try:
                # Extraire les données depuis le fichier
                student_id = row[col_map['student_id']]
                matiere = row[col_map['matiere']]
                type_devoir = row[col_map['type_devoir']]
                trimestre = row[col_map['trimestre']]
                note = row[col_map['note']]
                coefficient = row[col_map['coefficient']]
                date_devoir = row[col_map['date_devoir']]
                appreciation = row[col_map.get('appreciation')] if 'appreciation' in col_map else None
                devoir_numero = row[col_map.get('devoir_numero')] if 'devoir_numero' in col_map else None
                classe_id = row[col_map.get('classe_id')] if 'classe_id' in col_map else None
                
                # Validation des données obligatoires
                if not student_id or not matiere or not type_devoir or trimestre is None or note is None or coefficient is None or not date_devoir:
                    results["errors"] += 1
                    results["error_details"].append(f"Ligne {row_num}: Données obligatoires manquantes")
                    continue
                
                # Validation de la note
                try:
                    note = float(note)
                    if note < 0 or note > 20:
                        results["errors"] += 1
                        results["error_details"].append(f"Ligne {row_num}: Note {note} invalide (doit être entre 0 et 20)")
                        continue
                except (ValueError, TypeError):
                    results["errors"] += 1
                    results["error_details"].append(f"Ligne {row_num}: Note '{note}' n'est pas un nombre valide")
                    continue
                
                # Validation du coefficient
                try:
                    coefficient = int(coefficient)
                    if coefficient < 1:
                        results["errors"] += 1
                        results["error_details"].append(f"Ligne {row_num}: Coefficient {coefficient} invalide (doit être >= 1)")
                        continue
                except (ValueError, TypeError):
                    results["errors"] += 1
                    results["error_details"].append(f"Ligne {row_num}: Coefficient '{coefficient}' n'est pas un nombre valide")
                    continue
                
                # Validation du trimestre
                try:
                    trimestre = int(trimestre)
                    if trimestre not in [1, 2, 3]:
                        results["errors"] += 1
                        results["error_details"].append(f"Ligne {row_num}: Trimestre {trimestre} invalide (doit être 1, 2 ou 3)")
                        continue
                except (ValueError, TypeError):
                    results["errors"] += 1
                    results["error_details"].append(f"Ligne {row_num}: Trimestre '{trimestre}' n'est pas un nombre valide")
                    continue
                
                # Validation de la date
                if isinstance(date_devoir, str):
                    try:
                        from datetime import datetime
                        date_devoir = datetime.strptime(date_devoir, "%Y-%m-%d").date()
                    except ValueError:
                        try:
                            # Essayer d'autres formats de date
                            date_devoir = datetime.strptime(date_devoir, "%d/%m/%Y").date()
                        except ValueError:
                            results["errors"] += 1
                            results["error_details"].append(f"Ligne {row_num}: Date '{date_devoir}' format invalide (attendu: YYYY-MM-DD ou DD/MM/YYYY)")
                            continue
                
                # Vérifier que l'élève existe
                student = db.query(models.Eleve).filter(models.Eleve.id == student_id).first()
                if not student:
                    results["errors"] += 1
                    results["error_details"].append(f"Ligne {row_num}: Élève ID {student_id} non trouvé")
                    continue
                
                # Utiliser la classe de l'élève si non spécifiée
                if not classe_id:
                    classe_id = student.classe_id
                
                # Vérifier que la classe existe
                db_class = db.query(models.Classe).filter(models.Classe.id == classe_id).first()
                if not db_class:
                    results["errors"] += 1
                    results["error_details"].append(f"Ligne {row_num}: Classe ID {classe_id} non trouvée")
                    continue
                
                # Créer l'évaluation
                evaluation = models.Evaluation(
                    matiere=str(matiere),
                    type=str(type_devoir),
                    trimestre=trimestre,
                    note=note,
                    coefficient=coefficient,
                    date=date_devoir,
                    appreciation=str(appreciation) if appreciation else None,
                    classe_id=classe_id,
                    eleve_id=student_id,
                    devoir_numero=str(devoir_numero) if devoir_numero else None,
                    valide=False
                )
                
                db.add(evaluation)
                results["success"] += 1
                
            except Exception as e:
                results["errors"] += 1
                results["error_details"].append(f"Ligne {row_num}: {str(e)}")
        
        db.commit()
        
        return {
            "message": f"Import terminé: {results['success']} notes importées, {results['errors']} erreurs",
            "results": results
        }
        
    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Erreur lors de l'import Excel: {str(e)}")
